"""
docx_report.py
--------------
Modul generasi laporan profesional format DOCX (Microsoft Word) untuk
platform Analisis Sentimen Media Sosial Indonesia.

Menggunakan python-docx untuk membuat laporan multi-halaman dengan
chart embedding, tabel, styling konsisten, dan struktur laporan lengkap.
"""
from __future__ import annotations

import io
import logging
from datetime import datetime
from typing import Any

import numpy as np
import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Mm, Pt, RGBColor

logger = logging.getLogger(__name__)

# ── Konstanta Styling ─────────────────────────────────────────────────────────
FONT_BODY = "Calibri"
FONT_HEADING = "Calibri Light"
COLOR_PRIMARY = RGBColor(0x1E, 0x3A, 0x5F)      # #1e3a5f - dark blue
COLOR_SECONDARY = RGBColor(0x0F, 0x17, 0x2A)     # #0f172a
COLOR_ACCENT = RGBColor(0x3B, 0x82, 0xF6)        # #3b82f6
COLOR_TEXT = RGBColor(0x1E, 0x29, 0x3B)           # #1e293b
COLOR_MUTED = RGBColor(0x64, 0x74, 0x8B)         # #64748b
COLOR_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
COLOR_TABLE_HEADER = RGBColor(0x1E, 0x3A, 0x5F)  # header row bg
COLOR_TABLE_ALT = RGBColor(0xF1, 0xF5, 0xF9)     # alternating row bg

A4_WIDTH_CM = 21.0
A4_HEIGHT_CM = 29.7
MARGIN_CM = 2.54


# ── Helper Functions ──────────────────────────────────────────────────────────

def _safe_get(data: dict, key: str, default: Any = None) -> Any:
    """Ambil nilai dari dict dengan aman."""
    if data is None:
        return default
    return data.get(key, default)


def _format_number(val: Any) -> str:
    """Format angka untuk tampilan."""
    if val is None or (isinstance(val, float) and np.isnan(val)):
        return "0"
    try:
        num = float(val)
        if num == int(num) and abs(num) < 1e15:
            return f"{int(num):,}".replace(",", ".")
        return f"{num:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    except (ValueError, TypeError):
        return str(val)


def _format_pct(val: Any) -> str:
    """Format persentase."""
    if val is None:
        return "0,0%"
    try:
        return f"{float(val):.1f}%".replace(".", ",")
    except (ValueError, TypeError):
        return str(val)


def _set_cell_shading(cell: Any, color_hex: str) -> None:
    """Set background shading untuk cell tabel."""
    shading_elm = OxmlElement("w:shd")
    shading_elm.set(qn("w:fill"), color_hex)
    shading_elm.set(qn("w:val"), "clear")
    cell._tc.get_or_add_tcPr().append(shading_elm)


def _add_page_number(section: Any) -> None:
    """Tambah nomor halaman di footer."""
    footer = section.footer
    footer.is_linked_to_previous = False
    paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = paragraph.add_run()
    run.font.size = Pt(9)
    run.font.color.rgb = COLOR_MUTED
    run.font.name = FONT_BODY

    # Field code untuk nomor halaman
    fldChar1 = OxmlElement("w:fldChar")
    fldChar1.set(qn("w:fldCharType"), "begin")
    run._r.append(fldChar1)

    instrText = OxmlElement("w:instrText")
    instrText.set(qn("xml:space"), "preserve")
    instrText.text = " PAGE "
    run._r.append(instrText)

    fldChar2 = OxmlElement("w:fldChar")
    fldChar2.set(qn("w:fldCharType"), "end")
    run._r.append(fldChar2)

    run2 = paragraph.add_run(" / ")
    run2.font.size = Pt(9)
    run2.font.color.rgb = COLOR_MUTED

    run3 = paragraph.add_run()
    run3.font.size = Pt(9)
    run3.font.color.rgb = COLOR_MUTED

    fldChar3 = OxmlElement("w:fldChar")
    fldChar3.set(qn("w:fldCharType"), "begin")
    run3._r.append(fldChar3)

    instrText2 = OxmlElement("w:instrText")
    instrText2.set(qn("xml:space"), "preserve")
    instrText2.text = " NUMPAGES "
    run3._r.append(instrText2)

    fldChar4 = OxmlElement("w:fldChar")
    fldChar4.set(qn("w:fldCharType"), "end")
    run3._r.append(fldChar4)


def _add_styled_heading(doc: Document, text: str, level: int = 1) -> None:
    """Tambah heading dengan styling kustom."""
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.name = FONT_HEADING
        run.font.color.rgb = COLOR_PRIMARY


def _add_paragraph(doc: Document, text: str, bold: bool = False,
                   italic: bool = False, font_size: int = 11,
                   alignment: int | None = None,
                   color: RGBColor | None = None) -> None:
    """Tambah paragraf dengan styling."""
    p = doc.add_paragraph()
    if alignment is not None:
        p.alignment = alignment
    run = p.add_run(text)
    run.font.name = FONT_BODY
    run.font.size = Pt(font_size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color


def _add_styled_table(doc: Document, headers: list[str],
                      rows: list[list[str]],
                      col_widths: list[float] | None = None) -> None:
    """
    Tambah tabel dengan header berwarna dan alternating rows.

    Parameters
    ----------
    doc : Document
        Dokumen DOCX.
    headers : list[str]
        Daftar header kolom.
    rows : list[list[str]]
        Data baris.
    col_widths : list[float] | None
        Lebar kolom dalam cm (opsional).
    """
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    # Header row
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(header)
        run.bold = True
        run.font.name = FONT_BODY
        run.font.size = Pt(10)
        run.font.color.rgb = COLOR_WHITE
        _set_cell_shading(cell, "1E3A5F")

    # Data rows
    for row_idx, row_data in enumerate(rows):
        for col_idx, cell_text in enumerate(row_data):
            if col_idx >= len(headers):
                break
            cell = table.rows[row_idx + 1].cells[col_idx]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(str(cell_text))
            run.font.name = FONT_BODY
            run.font.size = Pt(9)
            run.font.color.rgb = COLOR_TEXT

            # Alternating row colors
            if row_idx % 2 == 1:
                _set_cell_shading(cell, "F1F5F9")

    # Set column widths jika disediakan
    if col_widths:
        for i, width in enumerate(col_widths):
            if i < len(headers):
                for row in table.rows:
                    row.cells[i].width = Cm(width)


def _add_chart_image(doc: Document, png_bytes: bytes,
                     width: float = 15.0) -> None:
    """Embed chart PNG ke dalam dokumen."""
    if png_bytes is None:
        _add_paragraph(doc, "[Chart tidak tersedia]", italic=True, color=COLOR_MUTED)
        return

    stream = io.BytesIO(png_bytes)
    doc.add_picture(stream, width=Cm(width))
    last_paragraph = doc.paragraphs[-1]
    last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _add_page_break(doc: Document) -> None:
    """Tambah page break."""
    doc.add_page_break()


def _add_toc_field(doc: Document) -> None:
    """Tambah field Table of Contents otomatis."""
    paragraph = doc.add_paragraph()
    run = paragraph.add_run()

    fldChar1 = OxmlElement("w:fldChar")
    fldChar1.set(qn("w:fldCharType"), "begin")
    run._r.append(fldChar1)

    instrText = OxmlElement("w:instrText")
    instrText.set(qn("xml:space"), "preserve")
    instrText.text = ' TOC \\o "1-3" \\h \\z \\u '
    run._r.append(instrText)

    fldChar2 = OxmlElement("w:fldChar")
    fldChar2.set(qn("w:fldCharType"), "separate")
    run._r.append(fldChar2)

    # Placeholder text
    run2 = paragraph.add_run(
        "Daftar isi akan diperbarui saat dokumen dibuka di Microsoft Word. "
        "Klik kanan → Update Field."
    )
    run2.font.color.rgb = COLOR_MUTED
    run2.font.size = Pt(10)
    run2.italic = True

    fldChar3 = OxmlElement("w:fldChar")
    fldChar3.set(qn("w:fldCharType"), "end")
    run2._r.append(fldChar3)


# ── Section Builders ──────────────────────────────────────────────────────────

def _build_cover_page(doc: Document, data: dict) -> None:
    """Halaman sampul laporan."""
    # Spacing sebelum judul
    for _ in range(6):
        doc.add_paragraph()

    # Judul utama
    _add_paragraph(doc, "LAPORAN ANALISIS SENTIMEN", bold=True, font_size=28,
                   alignment=WD_ALIGN_PARAGRAPH.CENTER, color=COLOR_PRIMARY)
    _add_paragraph(doc, "MEDIA SOSIAL", bold=True, font_size=28,
                   alignment=WD_ALIGN_PARAGRAPH.CENTER, color=COLOR_PRIMARY)

    doc.add_paragraph()

    # Garis pemisah
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("━" * 40)
    run.font.color.rgb = COLOR_ACCENT
    run.font.size = Pt(14)

    doc.add_paragraph()

    # Subtitle
    date_range = _safe_get(data, "date_range", "")
    if date_range:
        _add_paragraph(doc, f"Periode: {date_range}", font_size=14,
                       alignment=WD_ALIGN_PARAGRAPH.CENTER, color=COLOR_TEXT)

    analysis_mode = _safe_get(data, "analysis_mode", "Full Analysis")
    _add_paragraph(doc, f"Mode Analisis: {analysis_mode}", font_size=12,
                   alignment=WD_ALIGN_PARAGRAPH.CENTER, color=COLOR_MUTED)

    doc.add_paragraph()
    doc.add_paragraph()

    # Informasi generasi
    now = datetime.now().strftime("%d %B %Y, %H:%M WIB")
    _add_paragraph(doc, f"Dibuat oleh: Pahamdata.com", font_size=11,
                   alignment=WD_ALIGN_PARAGRAPH.CENTER, color=COLOR_MUTED)
    _add_paragraph(doc, f"Tanggal: {now}", font_size=11,
                   alignment=WD_ALIGN_PARAGRAPH.CENTER, color=COLOR_MUTED)

    session_id = _safe_get(data, "session_id", "")
    if session_id:
        _add_paragraph(doc, f"Session ID: {session_id}", font_size=9,
                       alignment=WD_ALIGN_PARAGRAPH.CENTER, color=COLOR_MUTED)


def _build_toc(doc: Document) -> None:
    """Halaman Daftar Isi."""
    _add_styled_heading(doc, "Daftar Isi", level=1)
    _add_toc_field(doc)


def _build_executive_summary(doc: Document, data: dict) -> None:
    """Bagian Executive Summary."""
    _add_styled_heading(doc, "Executive Summary", level=1)

    summary_text = _safe_get(data, "summary_text", "")
    if summary_text:
        paragraphs = summary_text.split("\n")
        for para in paragraphs:
            para = para.strip()
            if para:
                _add_paragraph(doc, para, font_size=11)
    else:
        _add_paragraph(doc, "Executive summary belum tersedia.", italic=True,
                       color=COLOR_MUTED)


def _build_methodology(doc: Document) -> None:
    """Bagian Metodologi."""
    _add_styled_heading(doc, "Metodologi", level=1)

    methodology_text = (
        "Analisis sentimen dalam laporan ini menggunakan model IndoBERT "
        "(Indonesian Bidirectional Encoder Representations from Transformers) "
        "yang telah di-fine-tune untuk tugas klasifikasi sentimen bahasa Indonesia. "
        "Model ini dikembangkan berdasarkan arsitektur BERT dan dilatih pada "
        "corpus bahasa Indonesia yang besar."
    )
    _add_paragraph(doc, methodology_text)

    _add_styled_heading(doc, "Tahapan Analisis", level=2)

    steps = [
        "Preprocessing: Pembersihan teks dari URL, mention, hashtag, emoji berlebih, "
        "dan normalisasi kata tidak baku (slang).",
        "Analisis Sentimen: Klasifikasi teks ke kategori Positif, Netral, atau Negatif "
        "menggunakan model IndoBERT fine-tuned.",
        "Analisis Emosi: Deteksi emosi (Marah, Senang, Sedih, Takut, Jijik, Terkejut, Netral) "
        "menggunakan model klasifikasi emosi.",
        "Deteksi Sarkasme: Identifikasi konten sarkastik yang dapat mempengaruhi "
        "interpretasi sentimen.",
        "Topic Modeling: Ekstraksi topik utama menggunakan BERTopic untuk memahami "
        "tema percakapan.",
        "Analisis Engagement: Perhitungan skor pengaruh berdasarkan metrik interaksi "
        "(likes, repost, komentar, views).",
    ]
    for i, step in enumerate(steps, 1):
        _add_paragraph(doc, f"{i}. {step}", font_size=10)


def _build_dataset_description(doc: Document, data: dict) -> None:
    """Bagian Deskripsi Dataset."""
    _add_styled_heading(doc, "Deskripsi Dataset", level=1)

    kpis = _safe_get(data, "kpis", {})
    df = _safe_get(data, "df")
    total_posts = kpis.get("total_posts", len(df) if df is not None else 0)
    total_accounts = kpis.get("total_accounts", 0)
    date_range = _safe_get(data, "date_range", "N/A")

    info_rows = [
        ["Total Postingan", _format_number(total_posts)],
        ["Total Akun Unik", _format_number(total_accounts)],
        ["Rentang Tanggal", str(date_range)],
        ["Sumber Data", "Media Sosial (X / Twitter)"],
        ["Format File", "CSV / Excel"],
    ]
    _add_styled_table(doc, ["Parameter", "Nilai"], info_rows, col_widths=[6, 10])


def _build_sentiment_analysis(doc: Document, data: dict) -> None:
    """Bagian Hasil Analisis Sentimen."""
    from reporting.chart_renderer import render_sentiment_donut, render_sentiment_timeline

    _add_styled_heading(doc, "Hasil Analisis Sentimen", level=1)

    kpis = _safe_get(data, "kpis", {})

    # KPI Table
    _add_styled_heading(doc, "Ringkasan KPI Sentimen", level=2)
    kpi_rows = [
        ["Total Postingan", _format_number(kpis.get("total_posts", 0))],
        ["Sentimen Positif", _format_pct(kpis.get("positive_pct", 0))],
        ["Sentimen Netral", _format_pct(kpis.get("neutral_pct", 0))],
        ["Sentimen Negatif", _format_pct(kpis.get("negative_pct", 0))],
        ["Confidence Rata-rata", _format_pct(kpis.get("avg_confidence", 0))],
    ]
    _add_styled_table(doc, ["Metrik", "Nilai"], kpi_rows, col_widths=[6, 10])
    doc.add_paragraph()

    # Sentiment donut chart
    _add_styled_heading(doc, "Distribusi Sentimen", level=2)
    try:
        donut_bytes = render_sentiment_donut(
            kpis.get("positive_pct", 33.3),
            kpis.get("neutral_pct", 33.3),
            kpis.get("negative_pct", 33.3),
        )
        _add_chart_image(doc, donut_bytes, width=12)
    except Exception as e:
        logger.error("Gagal render sentiment donut: %s", e)
        _add_paragraph(doc, "[Chart sentimen gagal di-render]", italic=True,
                       color=COLOR_MUTED)

    # Sentiment timeline
    _add_styled_heading(doc, "Tren Sentimen Harian", level=2)
    daily = _safe_get(data, "daily")
    try:
        timeline_bytes = render_sentiment_timeline(daily)
        _add_chart_image(doc, timeline_bytes)
    except Exception as e:
        logger.error("Gagal render sentiment timeline: %s", e)
        _add_paragraph(doc, "[Chart timeline gagal di-render]", italic=True,
                       color=COLOR_MUTED)

    # Sentiment index
    sentiment_index = _safe_get(data, "sentiment_index")
    if sentiment_index is not None:
        doc.add_paragraph()
        _add_paragraph(doc, f"Indeks Sentimen: {sentiment_index:+.2f} "
                       f"(skala -100 hingga +100)", bold=True, font_size=12,
                       color=COLOR_PRIMARY)


def _build_emotion_analysis(doc: Document, data: dict) -> None:
    """Bagian Hasil Analisis Emosi."""
    from reporting.chart_renderer import render_emotion_bar

    _add_styled_heading(doc, "Hasil Analisis Emosi", level=1)

    emotion_dist = _safe_get(data, "emotion_distribution", {})

    # Tabel distribusi emosi
    if emotion_dist:
        total_emotions = sum(emotion_dist.values()) or 1
        emotion_rows = []
        for emotion, count in sorted(emotion_dist.items(),
                                     key=lambda x: x[1], reverse=True):
            pct = (count / total_emotions) * 100
            emotion_rows.append([emotion, _format_number(count), _format_pct(pct)])

        _add_styled_table(doc, ["Emosi", "Jumlah", "Persentase"], emotion_rows,
                          col_widths=[5, 4, 4])
        doc.add_paragraph()

    # Chart
    try:
        bar_bytes = render_emotion_bar(emotion_dist)
        _add_chart_image(doc, bar_bytes)
    except Exception as e:
        logger.error("Gagal render emotion bar: %s", e)
        _add_paragraph(doc, "[Chart emosi gagal di-render]", italic=True,
                       color=COLOR_MUTED)


def _build_sarcasm_analysis(doc: Document, data: dict) -> None:
    """Bagian Hasil Deteksi Sarkasme."""
    from reporting.chart_renderer import render_sarcasm_pie

    _add_styled_heading(doc, "Hasil Deteksi Sarkasme", level=1)

    sarcasm_dist = _safe_get(data, "sarcasm_distribution", {})
    sarcasm_count = sarcasm_dist.get("sarcasm", sarcasm_dist.get("Sarcasm", 0))
    non_sarcasm_count = sarcasm_dist.get("non_sarcasm",
                                        sarcasm_dist.get("Not Sarcasm", 0))
    total = sarcasm_count + non_sarcasm_count

    if total > 0:
        sarcasm_pct = (sarcasm_count / total) * 100
        rows = [
            ["Sarkasme", _format_number(sarcasm_count), _format_pct(sarcasm_pct)],
            ["Non-Sarkasme", _format_number(non_sarcasm_count),
             _format_pct(100 - sarcasm_pct)],
            ["Total", _format_number(total), "100,0%"],
        ]
        _add_styled_table(doc, ["Kategori", "Jumlah", "Persentase"], rows,
                          col_widths=[5, 4, 4])
        doc.add_paragraph()

    # Pie chart
    try:
        pie_bytes = render_sarcasm_pie(sarcasm_count, non_sarcasm_count)
        _add_chart_image(doc, pie_bytes, width=12)
    except Exception as e:
        logger.error("Gagal render sarcasm pie: %s", e)

    # Contoh postingan sarkastik
    df = _safe_get(data, "df")
    if df is not None and "sarcasm" in df.columns:
        sarcastic_posts = df[df["sarcasm"] == "Sarcasm"]
        if len(sarcastic_posts) > 0:
            doc.add_paragraph()
            _add_styled_heading(doc, "Contoh Postingan Sarkastik", level=2)
            sample = sarcastic_posts.head(5)
            for _, row in sample.iterrows():
                konten = str(row.get("Konten", row.get("clean_text", "")))[:200]
                akun = str(row.get("X akun", "N/A"))
                _add_paragraph(doc, f"@{akun}: \"{konten}\"", font_size=10,
                               italic=True, color=COLOR_MUTED)


def _build_topic_analysis(doc: Document, data: dict) -> None:
    """Bagian Topik Utama Percakapan."""
    from reporting.chart_renderer import render_topic_bar

    _add_styled_heading(doc, "Topik Utama Percakapan", level=1)

    topics = _safe_get(data, "topics", [])

    if topics:
        topic_rows = []
        for i, t in enumerate(topics[:15], 1):
            name = t.get("topic_name", f"Topik {i}")
            freq = _format_number(t.get("frequency", 0))
            keywords = t.get("keywords", t.get("representative_posts", ""))
            if isinstance(keywords, list):
                keywords = ", ".join(str(k) for k in keywords[:5])
            topic_rows.append([str(i), name, freq, str(keywords)[:80]])

        _add_styled_table(
            doc,
            ["No", "Topik", "Frekuensi", "Kata Kunci"],
            topic_rows,
            col_widths=[1.5, 4, 3, 7.5],
        )
        doc.add_paragraph()

        try:
            topic_chart = render_topic_bar(topics)
            _add_chart_image(doc, topic_chart)
        except Exception as e:
            logger.error("Gagal render topic bar: %s", e)
    else:
        _add_paragraph(doc, "Tidak ada data topik tersedia.", italic=True,
                       color=COLOR_MUTED)


def _build_influencer_analysis(doc: Document, data: dict) -> None:
    """Bagian Analisis Akun Berpengaruh."""
    from reporting.chart_renderer import render_engagement_bar

    _add_styled_heading(doc, "Analisis Akun Berpengaruh", level=1)

    accounts = _safe_get(data, "accounts")

    if accounts is not None:
        if isinstance(accounts, list):
            accounts_df = pd.DataFrame(accounts)
        else:
            accounts_df = accounts.copy()

        if len(accounts_df) > 0:
            # Identifikasi kolom
            akun_col = next((c for c in ["akun", "X akun", "account"] if c in accounts_df.columns),
                            accounts_df.columns[0])
            eng_col = next((c for c in ["engagement", "engagement_score", "influence_score"]
                            if c in accounts_df.columns), None)
            sent_col = next((c for c in ["dominant_sentiment", "sentiment"]
                             if c in accounts_df.columns), None)

            top_10 = accounts_df.nlargest(10, eng_col) if eng_col else accounts_df.head(10)

            rows = []
            for i, (_, row) in enumerate(top_10.iterrows(), 1):
                row_data = [str(i), str(row.get(akun_col, "N/A"))]
                if "total_post" in accounts_df.columns:
                    row_data.append(_format_number(row.get("total_post", 0)))
                if eng_col:
                    row_data.append(_format_number(row.get(eng_col, 0)))
                if sent_col:
                    row_data.append(str(row.get(sent_col, "N/A")))
                rows.append(row_data)

            headers = ["No", "Akun"]
            if "total_post" in accounts_df.columns:
                headers.append("Total Post")
            if eng_col:
                headers.append("Engagement")
            if sent_col:
                headers.append("Sentimen Dominan")

            _add_styled_table(doc, headers, rows)
            doc.add_paragraph()

            try:
                eng_chart = render_engagement_bar(accounts_df)
                _add_chart_image(doc, eng_chart)
            except Exception as e:
                logger.error("Gagal render engagement bar: %s", e)
        else:
            _add_paragraph(doc, "Tidak ada data akun.", italic=True, color=COLOR_MUTED)
    else:
        _add_paragraph(doc, "Data akun tidak tersedia.", italic=True, color=COLOR_MUTED)


def _build_engagement_analysis(doc: Document, data: dict) -> None:
    """Bagian Analisis Engagement."""
    _add_styled_heading(doc, "Analisis Engagement", level=1)

    kpis = _safe_get(data, "kpis", {})
    weighted = _safe_get(data, "weighted_sentiment", {})

    rows = [
        ["Total Likes", _format_number(kpis.get("total_likes", 0))],
        ["Total Repost", _format_number(kpis.get("total_repost", 0))],
        ["Total Komentar", _format_number(kpis.get("total_komentar", 0))],
        ["Total Views", _format_number(kpis.get("total_views", 0))],
        ["Rata-rata Engagement per Post",
         _format_number(kpis.get("avg_engagement", 0))],
    ]

    if weighted:
        rows.append(["Weighted Sentiment Score",
                      _format_number(weighted.get("score", 0))])
        rows.append(["Weighted Sentiment Label",
                      str(weighted.get("label", "N/A"))])

    _add_styled_table(doc, ["Metrik", "Nilai"], rows, col_widths=[8, 8])


def _build_daily_trend(doc: Document, data: dict) -> None:
    """Bagian Trend Harian."""
    _add_styled_heading(doc, "Tren Harian", level=1)

    daily = _safe_get(data, "daily")
    if daily is None:
        _add_paragraph(doc, "Data tren harian tidak tersedia.", italic=True,
                       color=COLOR_MUTED)
        return

    if isinstance(daily, list):
        daily_df = pd.DataFrame(daily)
    else:
        daily_df = daily.copy()

    if len(daily_df) == 0:
        _add_paragraph(doc, "Data tren harian kosong.", italic=True, color=COLOR_MUTED)
        return

    # Tabel (max 30 baris)
    display_df = daily_df.head(30)
    headers = ["Tanggal", "Total", "Positif", "Netral", "Negatif"]
    rows = []
    for _, row in display_df.iterrows():
        rows.append([
            str(row.get("date", ""))[:10],
            _format_number(row.get("total_post", row.get("total", 0))),
            _format_number(row.get("positive", 0)),
            _format_number(row.get("neutral", 0)),
            _format_number(row.get("negative", 0)),
        ])
    _add_styled_table(doc, headers, rows)

    if len(daily_df) > 30:
        _add_paragraph(doc, f"(Menampilkan 30 dari {len(daily_df)} hari)",
                       italic=True, font_size=9, color=COLOR_MUTED)


def _build_weekly_trend(doc: Document, data: dict) -> None:
    """Bagian Trend Mingguan (agregasi dari daily)."""
    _add_styled_heading(doc, "Tren Mingguan", level=1)

    daily = _safe_get(data, "daily")
    if daily is None:
        _add_paragraph(doc, "Data tidak tersedia untuk agregasi mingguan.",
                       italic=True, color=COLOR_MUTED)
        return

    if isinstance(daily, list):
        df = pd.DataFrame(daily)
    else:
        df = daily.copy()

    if len(df) == 0 or "date" not in df.columns:
        _add_paragraph(doc, "Data tren mingguan kosong.", italic=True,
                       color=COLOR_MUTED)
        return

    df["date"] = pd.to_datetime(df["date"], errors="coerce")
    df["week"] = df["date"].dt.isocalendar().week.astype(int)
    df["year"] = df["date"].dt.year

    num_cols = ["positive", "neutral", "negative"]
    existing_cols = [c for c in num_cols if c in df.columns]

    if not existing_cols:
        _add_paragraph(doc, "Kolom sentimen tidak ditemukan.", italic=True,
                       color=COLOR_MUTED)
        return

    weekly = df.groupby(["year", "week"])[existing_cols].sum().reset_index()
    weekly["total"] = weekly[existing_cols].sum(axis=1)
    weekly = weekly.sort_values(["year", "week"])

    headers = ["Minggu", "Total", "Positif", "Netral", "Negatif"]
    rows = []
    for _, row in weekly.iterrows():
        rows.append([
            f"W{int(row['week'])} ({int(row['year'])})",
            _format_number(row.get("total", 0)),
            _format_number(row.get("positive", 0)),
            _format_number(row.get("neutral", 0)),
            _format_number(row.get("negative", 0)),
        ])
    _add_styled_table(doc, headers, rows)


def _build_monthly_trend(doc: Document, data: dict) -> None:
    """Bagian Trend Bulanan (agregasi dari daily)."""
    _add_styled_heading(doc, "Tren Bulanan", level=1)

    daily = _safe_get(data, "daily")
    if daily is None:
        _add_paragraph(doc, "Data tidak tersedia untuk agregasi bulanan.",
                       italic=True, color=COLOR_MUTED)
        return

    if isinstance(daily, list):
        df = pd.DataFrame(daily)
    else:
        df = daily.copy()

    if len(df) == 0 or "date" not in df.columns:
        _add_paragraph(doc, "Data tren bulanan kosong.", italic=True,
                       color=COLOR_MUTED)
        return

    df["date"] = pd.to_datetime(df["date"], errors="coerce")
    df["month"] = df["date"].dt.to_period("M")

    num_cols = ["positive", "neutral", "negative"]
    existing_cols = [c for c in num_cols if c in df.columns]

    if not existing_cols:
        _add_paragraph(doc, "Kolom sentimen tidak ditemukan.", italic=True,
                       color=COLOR_MUTED)
        return

    monthly = df.groupby("month")[existing_cols].sum().reset_index()
    monthly["total"] = monthly[existing_cols].sum(axis=1)

    headers = ["Bulan", "Total", "Positif", "Netral", "Negatif"]
    rows = []
    for _, row in monthly.iterrows():
        rows.append([
            str(row["month"]),
            _format_number(row.get("total", 0)),
            _format_number(row.get("positive", 0)),
            _format_number(row.get("neutral", 0)),
            _format_number(row.get("negative", 0)),
        ])
    _add_styled_table(doc, headers, rows)


def _build_key_findings(doc: Document, data: dict) -> None:
    """Bagian Temuan Utama."""
    _add_styled_heading(doc, "Temuan Utama", level=1)

    kpis = _safe_get(data, "kpis", {})
    sentiment_index = _safe_get(data, "sentiment_index", 0)

    findings = []

    # Auto-generate findings dari KPI
    pos_pct = kpis.get("positive_pct", 0)
    neg_pct = kpis.get("negative_pct", 0)
    neu_pct = kpis.get("neutral_pct", 0)

    dominant = "Positif" if pos_pct >= neg_pct and pos_pct >= neu_pct else \
               "Negatif" if neg_pct >= pos_pct and neg_pct >= neu_pct else "Netral"
    findings.append(
        f"Sentimen dominan adalah {dominant} dengan persentase "
        f"Positif {_format_pct(pos_pct)}, Netral {_format_pct(neu_pct)}, "
        f"Negatif {_format_pct(neg_pct)}."
    )

    if sentiment_index is not None:
        try:
            idx = float(sentiment_index)
            tone = "sangat positif" if idx > 50 else "cenderung positif" if idx > 0 \
                else "cenderung negatif" if idx > -50 else "sangat negatif"
            findings.append(
                f"Indeks sentimen bernilai {idx:+.2f}, menunjukkan persepsi publik yang {tone}."
            )
        except (ValueError, TypeError):
            pass

    sarcasm_dist = _safe_get(data, "sarcasm_distribution", {})
    sarcasm_count = sarcasm_dist.get("sarcasm", sarcasm_dist.get("Sarcasm", 0))
    total_posts = kpis.get("total_posts", 0)
    if total_posts > 0 and sarcasm_count > 0:
        sarcasm_pct = (sarcasm_count / total_posts) * 100
        findings.append(
            f"Terdeteksi {_format_number(sarcasm_count)} postingan sarkastik "
            f"({_format_pct(sarcasm_pct)} dari total), yang perlu dipertimbangkan "
            f"dalam interpretasi sentimen."
        )

    # Sentiment spikes
    spikes = _safe_get(data, "sentiment_spikes", [])
    if spikes:
        findings.append(
            f"Ditemukan {len(spikes)} lonjakan sentimen signifikan selama periode analisis."
        )

    for i, finding in enumerate(findings, 1):
        _add_paragraph(doc, f"{i}. {finding}", font_size=11)


def _build_strategic_insights(doc: Document, data: dict) -> None:
    """Bagian Insight Strategis."""
    _add_styled_heading(doc, "Insight Strategis", level=1)

    kpis = _safe_get(data, "kpis", {})
    pos_pct = kpis.get("positive_pct", 0)
    neg_pct = kpis.get("negative_pct", 0)

    insights = [
        "Lakukan monitoring sentimen secara berkala untuk mengidentifikasi "
        "perubahan persepsi publik secara real-time.",
        "Fokus respons pada postingan dengan engagement tinggi untuk "
        "memaksimalkan dampak komunikasi.",
        "Perhatikan konten sarkastik yang mungkin menyamarkan sentimen negatif "
        "dalam bentuk humor atau sindiran.",
    ]

    if neg_pct > 30:
        insights.append(
            "Tingkat sentimen negatif cukup tinggi. Diperlukan strategi komunikasi "
            "krisis dan respons proaktif terhadap keluhan publik."
        )
    if pos_pct > 50:
        insights.append(
            "Sentimen positif dominan. Manfaatkan momentum ini untuk memperkuat "
            "brand image dan meningkatkan engagement."
        )

    for i, insight in enumerate(insights, 1):
        _add_paragraph(doc, f"{i}. {insight}", font_size=11)


def _build_conclusion(doc: Document, data: dict) -> None:
    """Bagian Kesimpulan."""
    _add_styled_heading(doc, "Kesimpulan", level=1)

    kpis = _safe_get(data, "kpis", {})
    total = kpis.get("total_posts", 0)
    pos_pct = kpis.get("positive_pct", 0)
    neg_pct = kpis.get("negative_pct", 0)

    conclusion = (
        f"Berdasarkan analisis terhadap {_format_number(total)} postingan media sosial, "
        f"diperoleh gambaran komprehensif mengenai persepsi publik. "
        f"Sentimen positif tercatat sebesar {_format_pct(pos_pct)} dan sentimen negatif "
        f"sebesar {_format_pct(neg_pct)}. "
        f"Hasil analisis ini dapat digunakan sebagai dasar pengambilan keputusan strategis "
        f"dalam pengelolaan komunikasi publik dan respons terhadap isu-isu yang berkembang."
    )
    _add_paragraph(doc, conclusion)


def _build_recommendations(doc: Document, data: dict) -> None:
    """Bagian Rekomendasi Kebijakan."""
    _add_styled_heading(doc, "Rekomendasi Kebijakan", level=1)

    recommendations = [
        "Implementasikan sistem monitoring sentimen real-time untuk deteksi dini "
        "isu negatif dan potensi krisis komunikasi.",
        "Buat tim respons cepat (rapid response team) untuk menangani lonjakan "
        "sentimen negatif dalam waktu kurang dari 24 jam.",
        "Kembangkan konten positif yang relevan dengan topik-topik utama yang "
        "teridentifikasi dalam analisis.",
        "Lakukan engagement aktif dengan akun-akun berpengaruh (influencer) untuk "
        "memperkuat narasi positif.",
        "Evaluasi efektivitas kebijakan komunikasi secara berkala menggunakan "
        "analisis sentimen sebagai salah satu indikator.",
        "Pertimbangkan penggunaan analisis sentimen untuk riset pasar dan "
        "pengembangan produk/layanan.",
    ]

    for i, rec in enumerate(recommendations, 1):
        _add_paragraph(doc, f"{i}. {rec}", font_size=11)


def _build_appendix(doc: Document, data: dict) -> None:
    """Bagian Lampiran Data (sampel 50 baris pertama)."""
    _add_styled_heading(doc, "Lampiran Data", level=1)

    df = _safe_get(data, "df")
    if df is None or len(df) == 0:
        _add_paragraph(doc, "Tidak ada data lampiran.", italic=True,
                       color=COLOR_MUTED)
        return

    _add_paragraph(doc, f"Menampilkan 50 baris pertama dari {len(df)} total data.",
                   italic=True, font_size=10, color=COLOR_MUTED)
    doc.add_paragraph()

    sample = df.head(50)
    display_cols = ["X akun", "Konten", "sentiment", "emotion", "sarcasm"]
    available_cols = [c for c in display_cols if c in sample.columns]

    if not available_cols:
        available_cols = list(sample.columns[:5])

    headers = available_cols
    rows = []
    for _, row in sample.iterrows():
        row_data = []
        for col in available_cols:
            val = str(row.get(col, ""))
            # Truncate long text untuk tabel
            if len(val) > 80:
                val = val[:77] + "..."
            row_data.append(val)
        rows.append(row_data)

    _add_styled_table(doc, headers, rows)


# ── Main API ──────────────────────────────────────────────────────────────────

def generate_docx_report(data: dict) -> bytes:
    """
    Generate laporan profesional multi-halaman format DOCX.

    Parameters
    ----------
    data : dict
        Dictionary berisi semua data analisis:
        - kpis: dict KPI sentimen
        - sentiment_index: float indeks sentimen
        - daily: list[dict] | DataFrame tren harian
        - accounts: list[dict] | DataFrame data akun
        - keywords: list[tuple] kata kunci
        - bigrams: list[tuple] bigrams
        - trigrams: list[tuple] trigrams
        - sentiment_spikes: list lonjakan sentimen
        - viral_posts: list | DataFrame posting viral
        - emotion_distribution: dict distribusi emosi
        - weighted_sentiment: dict sentimen berbobot
        - sarcasm_distribution: dict distribusi sarkasme
        - topics: list[dict] topik
        - summary_text: str executive summary
        - df: pd.DataFrame data lengkap
        - session_id: str ID sesi
        - analysis_mode: str mode analisis
        - date_range: str rentang tanggal

    Returns
    -------
    bytes
        DOCX file sebagai bytes.

    Raises
    ------
    Exception
        Jika gagal membuat dokumen (ditangani dengan fallback minimal).
    """
    if data is None:
        data = {}

    logger.info("Memulai generasi laporan DOCX...")

    try:
        doc = Document()

        # ── Page setup (A4) ───────────────────────────────────────────────
        section = doc.sections[0]
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.left_margin = Cm(MARGIN_CM)
        section.right_margin = Cm(MARGIN_CM)
        section.top_margin = Cm(MARGIN_CM)
        section.bottom_margin = Cm(MARGIN_CM)

        # Page numbers di footer
        _add_page_number(section)

        # ── Build sections ────────────────────────────────────────────────
        _build_cover_page(doc, data)
        _add_page_break(doc)

        _build_toc(doc)
        _add_page_break(doc)

        _build_executive_summary(doc, data)
        _add_page_break(doc)

        _build_methodology(doc)
        _add_page_break(doc)

        _build_dataset_description(doc, data)
        _add_page_break(doc)

        _build_sentiment_analysis(doc, data)
        _add_page_break(doc)

        _build_emotion_analysis(doc, data)
        _add_page_break(doc)

        _build_sarcasm_analysis(doc, data)
        _add_page_break(doc)

        _build_topic_analysis(doc, data)
        _add_page_break(doc)

        _build_influencer_analysis(doc, data)
        _add_page_break(doc)

        _build_engagement_analysis(doc, data)
        _add_page_break(doc)

        _build_daily_trend(doc, data)
        _add_page_break(doc)

        _build_weekly_trend(doc, data)
        _add_page_break(doc)

        _build_monthly_trend(doc, data)
        _add_page_break(doc)

        _build_key_findings(doc, data)
        _add_page_break(doc)

        _build_strategic_insights(doc, data)
        _add_page_break(doc)

        _build_conclusion(doc, data)
        _add_page_break(doc)

        _build_recommendations(doc, data)
        _add_page_break(doc)

        _build_appendix(doc, data)

        # ── Save ke bytes ─────────────────────────────────────────────────
        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)
        docx_bytes = buf.read()

        logger.info("Laporan DOCX berhasil digenerate (%d bytes).", len(docx_bytes))
        return docx_bytes

    except Exception as e:
        logger.error("Gagal generate DOCX report: %s", e, exc_info=True)
        # Fallback: dokumen minimal
        doc = Document()
        doc.add_heading("Laporan Analisis Sentimen", level=0)
        doc.add_paragraph(f"Gagal membuat laporan lengkap: {e}")
        doc.add_paragraph("Silakan coba lagi atau hubungi administrator.")
        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)
        return buf.read()
