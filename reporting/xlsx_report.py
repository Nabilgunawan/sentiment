"""
xlsx_report.py
--------------
Modul generasi laporan multi-sheet Excel (XLSX) untuk platform
Analisis Sentimen Media Sosial Indonesia.

Menggunakan openpyxl untuk styling profesional: header berwarna,
alternating rows, auto-adjust column widths, freeze panes, dll.
"""
from __future__ import annotations

import io
import logging
from typing import Any

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    NamedStyle,
    PatternFill,
    Side,
    numbers,
)
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

logger = logging.getLogger(__name__)

# ── Konstanta Styling ─────────────────────────────────────────────────────────
HEADER_FILL = PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)

ALT_ROW_FILL = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
NORMAL_FILL = PatternFill(fill_type=None)

DATA_FONT = Font(name="Calibri", size=10, color="1E293B")
DATA_ALIGNMENT = Alignment(vertical="top", wrap_text=True)
DATA_ALIGNMENT_CENTER = Alignment(horizontal="center", vertical="center")
DATA_ALIGNMENT_NUMBER = Alignment(horizontal="right", vertical="center")

THIN_BORDER = Border(
    left=Side(style="thin", color="CBD5E1"),
    right=Side(style="thin", color="CBD5E1"),
    top=Side(style="thin", color="CBD5E1"),
    bottom=Side(style="thin", color="CBD5E1"),
)

# Tab colors (hex tanpa #)
TAB_COLORS: dict[str, str] = {
    "Raw Data": "3B82F6",
    "Clean Data": "22C55E",
    "Sentiment Result": "EF4444",
    "Emotion Result": "A855F7",
    "Sarcasm Result": "F97316",
    "Topic Analysis": "06B6D4",
    "Daily Trend": "EAB308",
    "Influencer Analysis": "EC4899",
    "Viral Posts": "F43F5E",
    "KPI Dashboard": "10B981",
    "Executive Summary": "6366F1",
}

MAX_COL_WIDTH = 50
MIN_COL_WIDTH = 8


# ── Helper Functions ──────────────────────────────────────────────────────────

def _safe_get(data: dict, key: str, default: Any = None) -> Any:
    """Ambil nilai dari dict dengan aman."""
    if data is None:
        return default
    return data.get(key, default)


def _sanitize_value(val: Any) -> Any:
    """Bersihkan value untuk Excel compatibility."""
    if val is None:
        return ""
    if isinstance(val, float) and (np.isnan(val) or np.isinf(val)):
        return ""
    if isinstance(val, (pd.Timestamp,)):
        return val.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(val, (np.integer,)):
        return int(val)
    if isinstance(val, (np.floating,)):
        return float(val) if np.isfinite(val) else ""
    if isinstance(val, (list, dict)):
        return str(val)
    return val


def _auto_adjust_column_widths(ws: Worksheet) -> None:
    """Auto-adjust lebar kolom berdasarkan konten."""
    for col_cells in ws.columns:
        max_length = 0
        col_letter = get_column_letter(col_cells[0].column)

        for cell in col_cells:
            try:
                cell_value = str(cell.value) if cell.value is not None else ""
                # Ambil baris pertama saja untuk multi-line
                first_line = cell_value.split("\n")[0]
                max_length = max(max_length, len(first_line))
            except (TypeError, AttributeError):
                pass

        adjusted_width = min(max(max_length + 3, MIN_COL_WIDTH), MAX_COL_WIDTH)
        ws.column_dimensions[col_letter].width = adjusted_width


def _apply_header_style(ws: Worksheet) -> None:
    """Terapkan styling ke header row (baris 1)."""
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGNMENT
        cell.border = THIN_BORDER


def _apply_data_styles(ws: Worksheet, start_row: int = 2) -> None:
    """Terapkan styling ke data rows dengan alternating colors."""
    for row_idx, row in enumerate(ws.iter_rows(min_row=start_row,
                                                max_row=ws.max_row,
                                                max_col=ws.max_column), start=0):
        fill = ALT_ROW_FILL if row_idx % 2 == 1 else NORMAL_FILL
        for cell in row:
            cell.font = DATA_FONT
            cell.fill = fill
            cell.border = THIN_BORDER
            cell.alignment = DATA_ALIGNMENT

            # Number formatting
            if isinstance(cell.value, float):
                if 0 <= cell.value <= 1:
                    cell.number_format = "0.00%"
                else:
                    cell.number_format = "#,##0.00"
                cell.alignment = DATA_ALIGNMENT_NUMBER
            elif isinstance(cell.value, int):
                cell.number_format = "#,##0"
                cell.alignment = DATA_ALIGNMENT_NUMBER


def _freeze_header(ws: Worksheet) -> None:
    """Freeze baris header."""
    ws.freeze_panes = "A2"


def _set_tab_color(ws: Worksheet, sheet_name: str) -> None:
    """Set warna tab sheet."""
    color = TAB_COLORS.get(sheet_name, "94A3B8")
    ws.sheet_properties.tabColor = color


def _write_dataframe_to_sheet(ws: Worksheet, df: pd.DataFrame,
                               columns: list[str] | None = None) -> None:
    """
    Tulis DataFrame ke worksheet dengan styling.

    Parameters
    ----------
    ws : Worksheet
        Target worksheet.
    df : pd.DataFrame
        Data yang akan ditulis.
    columns : list[str] | None
        Kolom yang akan ditulis (default: semua kolom di df).
    """
    if df is None or len(df) == 0:
        ws.append(["Tidak ada data"])
        return

    if columns:
        available_cols = [c for c in columns if c in df.columns]
        if not available_cols:
            available_cols = list(df.columns)
    else:
        available_cols = list(df.columns)

    # Header
    ws.append(available_cols)

    # Data rows
    for _, row in df.iterrows():
        row_data = []
        for col in available_cols:
            val = _sanitize_value(row.get(col, ""))
            # Truncate teks sangat panjang
            if isinstance(val, str) and len(val) > 2000:
                val = val[:1997] + "..."
            row_data.append(val)
        ws.append(row_data)

    # Apply styles
    _apply_header_style(ws)
    _apply_data_styles(ws)
    _freeze_header(ws)
    _auto_adjust_column_widths(ws)


# ── Sheet Builders ────────────────────────────────────────────────────────────

def _build_raw_data_sheet(wb: Workbook, data: dict) -> None:
    """Sheet 1: Raw Data."""
    ws = wb.active
    ws.title = "Raw Data"
    _set_tab_color(ws, "Raw Data")

    df = _safe_get(data, "df")
    if df is None:
        ws.append(["Tidak ada data"])
        return

    columns = ["Tanggal", "Waktu", "X akun", "Konten", "Komentar",
                "Repost", "Likes", "Views", "Link"]
    _write_dataframe_to_sheet(ws, df, columns)
    logger.debug("Sheet 'Raw Data' berhasil dibuat (%d baris).", len(df))


def _build_clean_data_sheet(wb: Workbook, data: dict) -> None:
    """Sheet 2: Clean Data."""
    ws = wb.create_sheet("Clean Data")
    _set_tab_color(ws, "Clean Data")

    df = _safe_get(data, "df")
    if df is None:
        ws.append(["Tidak ada data"])
        return

    columns = ["Konten", "clean_text"]
    _write_dataframe_to_sheet(ws, df, columns)
    logger.debug("Sheet 'Clean Data' berhasil dibuat.")


def _build_sentiment_sheet(wb: Workbook, data: dict) -> None:
    """Sheet 3: Sentiment Result."""
    ws = wb.create_sheet("Sentiment Result")
    _set_tab_color(ws, "Sentiment Result")

    df = _safe_get(data, "df")
    if df is None:
        ws.append(["Tidak ada data"])
        return

    columns = ["X akun", "Konten", "sentiment", "confidence",
                "prob_positive", "prob_neutral", "prob_negative"]
    _write_dataframe_to_sheet(ws, df, columns)
    logger.debug("Sheet 'Sentiment Result' berhasil dibuat.")


def _build_emotion_sheet(wb: Workbook, data: dict) -> None:
    """Sheet 4: Emotion Result."""
    ws = wb.create_sheet("Emotion Result")
    _set_tab_color(ws, "Emotion Result")

    df = _safe_get(data, "df")
    if df is None:
        ws.append(["Tidak ada data"])
        return

    columns = ["X akun", "Konten", "emotion", "confidence"]
    # Jika ada kolom confidence terpisah untuk emotion
    if "emotion_confidence" in df.columns:
        columns = ["X akun", "Konten", "emotion", "emotion_confidence"]
    _write_dataframe_to_sheet(ws, df, columns)
    logger.debug("Sheet 'Emotion Result' berhasil dibuat.")


def _build_sarcasm_sheet(wb: Workbook, data: dict) -> None:
    """Sheet 5: Sarcasm Result."""
    ws = wb.create_sheet("Sarcasm Result")
    _set_tab_color(ws, "Sarcasm Result")

    df = _safe_get(data, "df")
    if df is None:
        ws.append(["Tidak ada data"])
        return

    columns = ["X akun", "Konten", "sarcasm", "confidence"]
    if "sarcasm_confidence" in df.columns:
        columns = ["X akun", "Konten", "sarcasm", "sarcasm_confidence"]
    _write_dataframe_to_sheet(ws, df, columns)
    logger.debug("Sheet 'Sarcasm Result' berhasil dibuat.")


def _build_topic_sheet(wb: Workbook, data: dict) -> None:
    """Sheet 6: Topic Analysis."""
    ws = wb.create_sheet("Topic Analysis")
    _set_tab_color(ws, "Topic Analysis")

    topics = _safe_get(data, "topics", [])
    if not topics:
        ws.append(["Tidak ada data topik"])
        return

    ws.append(["topic_name", "frequency", "representative_posts"])
    _apply_header_style(ws)

    for i, topic in enumerate(topics):
        name = topic.get("topic_name", f"Topik {i + 1}")
        freq = topic.get("frequency", 0)

        posts = topic.get("representative_posts", "")
        if isinstance(posts, list):
            posts = "\n".join(str(p) for p in posts[:5])
        elif not isinstance(posts, str):
            posts = str(posts)

        ws.append([_sanitize_value(name), _sanitize_value(freq),
                   _sanitize_value(posts)])

    _apply_data_styles(ws)
    _freeze_header(ws)
    _auto_adjust_column_widths(ws)
    logger.debug("Sheet 'Topic Analysis' berhasil dibuat (%d topik).", len(topics))


def _build_daily_trend_sheet(wb: Workbook, data: dict) -> None:
    """Sheet 7: Daily Trend."""
    ws = wb.create_sheet("Daily Trend")
    _set_tab_color(ws, "Daily Trend")

    daily = _safe_get(data, "daily")
    if daily is None:
        ws.append(["Tidak ada data tren harian"])
        return

    if isinstance(daily, list):
        daily_df = pd.DataFrame(daily)
    else:
        daily_df = daily.copy()

    columns = ["date", "total_post", "positive", "neutral", "negative", "engagement"]
    _write_dataframe_to_sheet(ws, daily_df, columns)
    logger.debug("Sheet 'Daily Trend' berhasil dibuat.")


def _build_influencer_sheet(wb: Workbook, data: dict) -> None:
    """Sheet 8: Influencer Analysis."""
    ws = wb.create_sheet("Influencer Analysis")
    _set_tab_color(ws, "Influencer Analysis")

    accounts = _safe_get(data, "accounts")
    if accounts is None:
        ws.append(["Tidak ada data akun"])
        return

    if isinstance(accounts, list):
        accounts_df = pd.DataFrame(accounts)
    else:
        accounts_df = accounts.copy()

    columns = ["akun", "total_post", "engagement", "dominant_sentiment",
                "influence_score"]
    # Fallback kolom
    if "akun" not in accounts_df.columns and "X akun" in accounts_df.columns:
        accounts_df = accounts_df.rename(columns={"X akun": "akun"})

    _write_dataframe_to_sheet(ws, accounts_df, columns)
    logger.debug("Sheet 'Influencer Analysis' berhasil dibuat.")


def _build_viral_posts_sheet(wb: Workbook, data: dict) -> None:
    """Sheet 9: Viral Posts."""
    ws = wb.create_sheet("Viral Posts")
    _set_tab_color(ws, "Viral Posts")

    viral = _safe_get(data, "viral_posts")
    if viral is None:
        # Fallback: ambil top posts dari df
        df = _safe_get(data, "df")
        if df is not None and len(df) > 0:
            # Coba sortir berdasarkan engagement
            eng_cols = ["engagement_score", "Likes", "Views"]
            sort_col = next((c for c in eng_cols if c in df.columns), None)
            if sort_col:
                viral = df.nlargest(50, sort_col)
            else:
                viral = df.head(50)
        else:
            ws.append(["Tidak ada data viral posts"])
            return

    if isinstance(viral, list):
        viral_df = pd.DataFrame(viral)
    else:
        viral_df = viral.copy()

    columns = ["X akun", "Konten", "sentiment", "emotion", "sarcasm",
                "engagement_score", "Likes", "Repost", "Komentar", "Views"]
    # Rename jika perlu
    rename_map = {"akun": "X akun", "konten": "Konten", "likes": "Likes",
                  "repost": "Repost", "komentar": "Komentar", "views": "Views"}
    viral_df = viral_df.rename(columns={k: v for k, v in rename_map.items()
                                        if k in viral_df.columns})

    _write_dataframe_to_sheet(ws, viral_df, columns)
    logger.debug("Sheet 'Viral Posts' berhasil dibuat.")


def _build_kpi_dashboard_sheet(wb: Workbook, data: dict) -> None:
    """Sheet 10: KPI Dashboard."""
    ws = wb.create_sheet("KPI Dashboard")
    _set_tab_color(ws, "KPI Dashboard")

    kpis = _safe_get(data, "kpis", {})
    sentiment_index = _safe_get(data, "sentiment_index")
    weighted = _safe_get(data, "weighted_sentiment", {})

    # Header
    ws.append(["KPI", "Nilai"])
    _apply_header_style(ws)

    # KPI entries
    kpi_entries = [
        ("Total Postingan", kpis.get("total_posts", 0)),
        ("Total Akun Unik", kpis.get("total_accounts", 0)),
        ("Sentimen Positif (%)", kpis.get("positive_pct", 0)),
        ("Sentimen Netral (%)", kpis.get("neutral_pct", 0)),
        ("Sentimen Negatif (%)", kpis.get("negative_pct", 0)),
        ("Jumlah Positif", kpis.get("positive_count", 0)),
        ("Jumlah Netral", kpis.get("neutral_count", 0)),
        ("Jumlah Negatif", kpis.get("negative_count", 0)),
        ("Confidence Rata-rata", kpis.get("avg_confidence", 0)),
        ("Total Likes", kpis.get("total_likes", 0)),
        ("Total Repost", kpis.get("total_repost", 0)),
        ("Total Komentar", kpis.get("total_komentar", 0)),
        ("Total Views", kpis.get("total_views", 0)),
        ("Rata-rata Engagement", kpis.get("avg_engagement", 0)),
    ]

    if sentiment_index is not None:
        kpi_entries.append(("Indeks Sentimen", sentiment_index))

    if weighted:
        kpi_entries.append(("Weighted Sentiment Score", weighted.get("score", 0)))
        kpi_entries.append(("Weighted Sentiment Label", weighted.get("label", "N/A")))

    sarcasm_dist = _safe_get(data, "sarcasm_distribution", {})
    if sarcasm_dist:
        sarcasm_count = sarcasm_dist.get("sarcasm", sarcasm_dist.get("Sarcasm", 0))
        non_sarcasm = sarcasm_dist.get("non_sarcasm", sarcasm_dist.get("Not Sarcasm", 0))
        kpi_entries.append(("Postingan Sarkastik", sarcasm_count))
        kpi_entries.append(("Postingan Non-Sarkastik", non_sarcasm))

    for kpi_name, kpi_val in kpi_entries:
        ws.append([kpi_name, _sanitize_value(kpi_val)])

    _apply_data_styles(ws)
    _freeze_header(ws)
    _auto_adjust_column_widths(ws)

    # Extra styling: KPI name column wider
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 20

    logger.debug("Sheet 'KPI Dashboard' berhasil dibuat (%d KPIs).", len(kpi_entries))


def _build_summary_sheet(wb: Workbook, data: dict) -> None:
    """Sheet 11: Executive Summary."""
    ws = wb.create_sheet("Executive Summary")
    _set_tab_color(ws, "Executive Summary")

    summary_text = _safe_get(data, "summary_text", "Executive summary belum tersedia.")

    # Title row
    ws.append(["EXECUTIVE SUMMARY"])
    title_cell = ws["A1"]
    title_cell.font = Font(name="Calibri", bold=True, size=16, color="1E3A5F")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")

    # Merge cells untuk judul
    ws.merge_cells("A1:F1")

    ws.append([])  # empty row

    # Summary content
    ws.append(["Ringkasan Analisis"])
    ws["A3"].font = Font(name="Calibri", bold=True, size=12, color="1E3A5F")

    ws.append([])

    # Tulis summary text dalam paragraf
    paragraphs = summary_text.split("\n") if summary_text else ["Tidak ada data."]
    for para in paragraphs:
        para = para.strip()
        if para:
            ws.append([para])

    # Merge summary cells untuk lebar penuh
    max_row = ws.max_row
    for row_idx in range(5, max_row + 1):
        ws.merge_cells(f"A{row_idx}:F{row_idx}")
        cell = ws[f"A{row_idx}"]
        cell.font = DATA_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    # Tambah info tambahan
    ws.append([])
    ws.append([])

    date_range = _safe_get(data, "date_range", "N/A")
    session_id = _safe_get(data, "session_id", "N/A")
    analysis_mode = _safe_get(data, "analysis_mode", "N/A")

    info_row = ws.max_row + 1
    ws.append(["Periode", str(date_range)])
    ws.append(["Mode Analisis", str(analysis_mode)])
    ws.append(["Session ID", str(session_id)])

    for row_idx in range(info_row, ws.max_row + 1):
        ws[f"A{row_idx}"].font = Font(name="Calibri", bold=True, size=10,
                                       color="64748B")
        ws[f"B{row_idx}"].font = Font(name="Calibri", size=10, color="1E293B")

    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 40

    logger.debug("Sheet 'Executive Summary' berhasil dibuat.")


# ── Main API ──────────────────────────────────────────────────────────────────

def generate_xlsx_report(data: dict) -> bytes:
    """
    Generate laporan multi-sheet Excel (XLSX) dengan styling profesional.

    Parameters
    ----------
    data : dict
        Dictionary berisi semua data analisis:
        - df: pd.DataFrame data lengkap
        - kpis: dict KPI sentimen
        - sentiment_index: float indeks sentimen
        - daily: list[dict] | DataFrame tren harian
        - accounts: list[dict] | DataFrame data akun
        - viral_posts: list[dict] | DataFrame posting viral
        - emotion_distribution: dict distribusi emosi
        - weighted_sentiment: dict sentimen berbobot
        - sarcasm_distribution: dict distribusi sarkasme
        - topics: list[dict] topik
        - summary_text: str executive summary
        - session_id: str ID sesi
        - analysis_mode: str mode analisis
        - date_range: str rentang tanggal

    Returns
    -------
    bytes
        XLSX file sebagai bytes.
    """
    if data is None:
        data = {}

    logger.info("Memulai generasi laporan XLSX...")

    try:
        wb = Workbook()

        # Build semua sheet
        _build_raw_data_sheet(wb, data)       # Sheet 1
        _build_clean_data_sheet(wb, data)     # Sheet 2
        _build_sentiment_sheet(wb, data)      # Sheet 3
        _build_emotion_sheet(wb, data)        # Sheet 4
        _build_sarcasm_sheet(wb, data)        # Sheet 5
        _build_topic_sheet(wb, data)          # Sheet 6
        _build_daily_trend_sheet(wb, data)    # Sheet 7
        _build_influencer_sheet(wb, data)     # Sheet 8
        _build_viral_posts_sheet(wb, data)    # Sheet 9
        _build_kpi_dashboard_sheet(wb, data)  # Sheet 10
        _build_summary_sheet(wb, data)        # Sheet 11

        # Save ke bytes
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        xlsx_bytes = buf.read()

        logger.info("Laporan XLSX berhasil digenerate (%d bytes, %d sheets).",
                     len(xlsx_bytes), len(wb.sheetnames))
        return xlsx_bytes

    except Exception as e:
        logger.error("Gagal generate XLSX report: %s", e, exc_info=True)
        # Fallback minimal
        wb = Workbook()
        ws = wb.active
        ws.title = "Error"
        ws.append(["Error generating report"])
        ws.append([str(e)])
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf.read()
